"""
엑셀 → 웹 데이터 자동 변환 스크립트
=====================================
사용법:
    python3 excel_to_data.py

엑셀 파일(검토상태표_v2.1.xlsx)의 '검토상태표' 시트에
새 행을 추가하고 이 스크립트를 실행하면 웹에 즉시 반영됩니다.

컬럼 순서 (엑셀 1행 헤더 기준):
  사례번호 | 감사기관 | 감사명 | 게시일 | 분야 | 지적제목 |
  요약 | 처분 | 관련근거 | 검토상태 | 검토메모 | 검토일 | 공식 게시물

공개 기준:
  검토상태가 'approved' 또는 'source_checked'인 행만 웹에 표시됩니다.
  'needs_review', 'rejected'는 제외됩니다.
"""

import hashlib
import json
import sys
from datetime import datetime, timezone, timedelta
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("오류: openpyxl 패키지가 없습니다. 아래 명령어로 설치하세요:")
    print("  pip install openpyxl")
    sys.exit(1)

# ── 경로 설정 ──────────────────────────────────────────────
HERE        = Path(__file__).resolve().parent
EXCEL_FILE  = HERE / "검토상태표_v2.1.xlsx"
SHEET_NAME  = "검토상태표"

# 출력할 폴더 목록 (필요에 따라 추가/제거)
OUTPUT_DIRS = [
    HERE / "barkley" / "깃허브_업로드용_v2.2 (1)" / "data",
    HERE / "barkley-test" / "data",
]

PUBLIC_STATUSES = {"approved", "source_checked"}

KST = timezone(timedelta(hours=9))

# ── 헤더 → JSON 키 매핑 ────────────────────────────────────
COL_MAP = {
    "사례번호":    "id",
    "감사기관":    "audit_org",
    "감사명":      "audit_name",
    "게시일":      "published_at",
    "분야":        "category",
    "지적제목":    "title",
    "요약":        "summary",
    "처분":        "disposition",
    "관련근거":    "legal_basis",
    "검토상태":    "review_status",
    "검토메모":    "review_note",
    "검토일":      "reviewed_at",
    "공식 게시물": "source_url",
}


def to_str(v):
    if v is None:
        return None
    if isinstance(v, datetime):
        return v.strftime("%Y-%m-%d")
    return str(v).strip() or None


def content_hash(case):
    raw = (case.get("title") or "") + "|" + (case.get("audit_name") or "")
    return hashlib.sha1(raw.encode()).hexdigest()[:12]


def parse_year(published_at):
    if not published_at:
        return None
    try:
        return int(str(published_at)[:4])
    except ValueError:
        return None


def parse_disposition_types(disposition):
    if not disposition:
        return []
    return [d.strip() for d in disposition.replace("/", "·").split("·") if d.strip()]


def read_excel(path, sheet):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        print(f"오류: '{sheet}' 시트를 찾을 수 없습니다.")
        print(f"  사용 가능한 시트: {wb.sheetnames}")
        sys.exit(1)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        print("오류: 시트가 비어 있습니다.")
        sys.exit(1)

    # 헤더 행
    raw_headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    headers = [COL_MAP.get(h, h) for h in raw_headers]

    cases = []
    skipped = 0
    for row in rows[1:]:
        if all(v is None for v in row):
            continue  # 빈 행 무시
        item = {headers[i]: to_str(v) for i, v in enumerate(row) if i < len(headers)}

        # 공개 여부 확인
        status = item.get("review_status") or ""
        if status not in PUBLIC_STATUSES:
            skipped += 1
            continue

        # 파생 필드
        item["audit_year"]         = parse_year(item.get("published_at"))
        item["disposition_types"]  = parse_disposition_types(item.get("disposition"))
        item["content_hash"]       = content_hash(item)
        item.setdefault("audit_type", None)
        item.setdefault("target_org", None)
        item.setdefault("source_url_scope", "page")
        item.setdefault("disposition_detail", None)

        # 관리자 전용 필드 제거
        for f in ("review_note", "reviewed_at"):
            item.pop(f, None)

        cases.append(item)

    return cases, skipped


def write_js(cases, out_dir, generated_at):
    data = {
        "schema_version": "2.2",
        "generated_at": generated_at,
        "counts": {"total": len(cases)},
        "cases": cases,
    }

    # JSON
    json_path = out_dir / "audit_cases.json"
    json_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")

    # JS (window.__AUDIT_DATA__ 형식)
    js_header = "/* 자동 생성 파일입니다. 직접 고치지 말고 excel_to_data.py 를 실행하세요. */\n"
    js_body   = f"window.__AUDIT_DATA__ = window.__AUDIT_DATA__ || {{}};\n" \
                f"window.__AUDIT_DATA__.cases = {json.dumps(data, ensure_ascii=False, indent=2)};\n"
    js_path = out_dir / "audit_cases.js"
    js_path.write_text(js_header + js_body, encoding="utf-8")

    print(f"  ✓ {json_path}")
    print(f"  ✓ {js_path}")


def main():
    if not EXCEL_FILE.exists():
        print(f"오류: 엑셀 파일을 찾을 수 없습니다.\n  {EXCEL_FILE}")
        sys.exit(1)

    print(f"엑셀 읽는 중: {EXCEL_FILE.name}")
    cases, skipped = read_excel(EXCEL_FILE, SHEET_NAME)
    generated_at = datetime.now(KST).strftime("%Y-%m-%dT%H:%M:%S+09:00")

    print(f"  공개 사례: {len(cases)}건  |  제외(검토대기·제외): {skipped}건")
    print()

    ok_count = 0
    for out_dir in OUTPUT_DIRS:
        if not out_dir.exists():
            print(f"  건너뜀 (폴더 없음): {out_dir}")
            continue
        print(f"저장 중: {out_dir}")
        write_js(cases, out_dir, generated_at)
        ok_count += 1

    print()
    if ok_count:
        print(f"완료! {ok_count}개 폴더에 저장됐습니다.")
        print("브라우저를 새로고침하면 변경사항이 반영됩니다.")
    else:
        print("경고: 저장된 폴더가 없습니다. OUTPUT_DIRS 경로를 확인하세요.")


if __name__ == "__main__":
    main()
