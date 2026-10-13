"""관리용 원본에서 직원용 검색 데이터를 만듭니다.

입력 (관리자만 다루는 파일)
  admin/cases_master.json                  감사사례 원본 (검토상태 포함)
  admin/synonyms_master.json               유의어 사전 원본
  admin/source_data/*자료원.xlsx           전국 감사자료원 엑셀

출력 (직원용 화면이 읽는 파일)
  dist/data/audit_cases.json  + .js        공개 가능한 사례만 (approved, source_checked)
  dist/data/sources.json      + .js        감사자료원 76개
  dist/data/synonyms.json     + .js        유의어 사전
  admin/review_status.xlsx                 검토상태표 (감사실 검토용)

.js 파일은 같은 내용을 <script> 로 읽을 수 있게 감싼 것입니다.
그래서 index.html 을 더블클릭해서 열어도(인터넷·서버 없이) 화면이 동작합니다.

실행:  python3 admin/build_data.py
오류가 있으면 아무 파일도 바꾸지 않고 멈춥니다.
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
from datetime import datetime, timezone, timedelta
from pathlib import Path

import openpyxl

ADMIN = Path(__file__).resolve().parent
ROOT = ADMIN.parent
DIST_DATA = ROOT / "dist" / "data"
KST = timezone(timedelta(hours=9))

REVIEW_STATUSES = {
    "needs_review": {"label": "검토대기", "public": False,
                      "desc": "자동추출 또는 일람표 수준으로 추가 검토 필요. 직원 화면에 내보내지 않음"},
    "source_checked": {"label": "출처확인", "public": True,
                        "desc": "출처와 제목은 확인했으나 법령·세부내용은 미확인. 주의표시 후 공개"},
    "approved": {"label": "검토완료", "public": True,
                 "desc": "감사실 검토자가 원문과 요약을 확인함"},
    "rejected": {"label": "제외", "public": False,
                 "desc": "출처불명·중복·오분류·비공개자료. 공개 금지"},
}
REQUIRED_CASE_FIELDS = ["id", "audit_org", "audit_year", "audit_name", "title", "source_url", "review_status"]
# 직원용 데이터에서 빼는 관리자 전용 항목
ADMIN_ONLY_FIELDS = {"attachment_url", "reviewed_by", "review_note", "legacy_ref"}

READINESS = {
    "ready": "수집 준비",
    "manual": "수동 확인",
    "recheck_required": "재확인 필요",
    "not_tested_environment": "환경 재시험",
}

errors: list[str] = []
warnings: list[str] = []


def now_kst() -> str:
    return datetime.now(KST).strftime("%Y-%m-%dT%H:%M:%S+09:00")


def norm_name(s: str) -> str:
    return re.sub(r"\s+", "", str(s or ""))


# ---------------------------------------------------------------- 감사사례
def build_cases() -> tuple[dict, list[dict]]:
    master = json.loads((ADMIN / "cases_master.json").read_text(encoding="utf-8"))
    cases = master["cases"]
    seen_ids, seen_hash = set(), {}
    public = []
    for c in cases:
        cid = c.get("id", "?")
        for f in REQUIRED_CASE_FIELDS:
            if c.get(f) in (None, ""):
                errors.append(f"[사례 {cid}] 필수항목 '{f}' 비어 있음")
        if cid in seen_ids:
            errors.append(f"[사례 {cid}] 사례번호 중복")
        seen_ids.add(cid)
        st = c.get("review_status")
        if st not in REVIEW_STATUSES:
            errors.append(f"[사례 {cid}] 알 수 없는 검토상태 '{st}'")
            continue
        url = str(c.get("source_url") or "")
        if not url.startswith("https://"):
            errors.append(f"[사례 {cid}] 공식 게시물 주소가 https:// 로 시작하지 않음: {url}")
        if st == "approved" and not c.get("summary"):
            errors.append(f"[사례 {cid}] 검토완료(approved)인데 요약이 비어 있음")
        h = hashlib.sha1("|".join([
            norm_name(c.get("audit_org")), norm_name(c.get("audit_name")), norm_name(c.get("title"))
        ]).encode("utf-8")).hexdigest()[:16]
        c["content_hash"] = h
        if st != "rejected":
            if h in seen_hash:
                errors.append(f"[사례 {cid}] 사례 {seen_hash[h]} 와 같은 지적사항(기관·감사명·제목 동일) — 중복")
            seen_hash[h] = cid
        if REVIEW_STATUSES[st]["public"]:
            item = {k: v for k, v in c.items() if k not in ADMIN_ONLY_FIELDS}
            item["disposition_types"] = [d.strip() for d in str(c.get("disposition") or "").split("·") if d.strip()]
            public.append(item)
    counts = {s: sum(1 for c in cases if c.get("review_status") == s) for s in REVIEW_STATUSES}
    out = {
        "schema_version": "2.0",
        "generated_at": now_kst(),
        "review_status_definitions": {k: {"label": v["label"], "desc": v["desc"]} for k, v in REVIEW_STATUSES.items()},
        "counts": {"master_total": len(cases), "public_total": len(public), "by_status": counts},
        "cases": public,
    }
    return out, cases


# ---------------------------------------------------------------- 감사자료원
def find_excel() -> Path:
    files = sorted((ADMIN / "source_data").glob("*.xlsx"))
    if not files:
        errors.append("admin/source_data 에 자료원 엑셀이 없음")
        raise SystemExit(report())
    return files[-1]


def source_group(category: str) -> str:
    c = category.strip()
    if c.startswith("중앙"):
        return "중앙"
    if c.startswith("광역"):
        return "광역"
    if c.startswith("서울"):
        return "서울자치구"
    return "경기도시군"


NO_DIRECT_LIST = ("검색", "경유", "메뉴경로", "부서", "필터")


def readiness_for(org: str, link_type: str, pilot: dict | None) -> tuple[str, str]:
    """수집준비도와 그 근거. 연찬 우선순위와는 별개입니다."""
    if pilot:
        verdict = pilot.get("final_verdict", "")
        if verdict.startswith("성공"):
            return "ready", "시범수집 검증 성공(목록·상세·첨부·PDF 추출)"
        if "부분" in verdict:
            return "not_tested_environment", "목록·상세·첨부주소 확인, 첨부 다운로드는 검증환경에서 502 — 다른 네트워크 재시험 필요"
        return "recheck_required", f"시범수집 검증 미완료: {verdict}"
    if link_type == "전국 통합검색":
        return "manual", "전국 통합검색 시스템(목록주소 확보) — 자동수집 대상 아님, 수동 검색용"
    if any(k in link_type for k in NO_DIRECT_LIST) and "사전정보공표" not in link_type:
        return "recheck_required", f"직접 목록주소 없음(링크 유형: {link_type}) — 실제 감사결과 목록 URL 확인 필요"
    return "manual", f"목록주소 확보(링크 유형: {link_type}), 게시판 구조는 미검증"


def build_sources() -> dict:
    path = find_excel()
    wb = openpyxl.load_workbook(path, data_only=False)

    # 연찬 우선기관 24곳
    priority = {}
    ws = wb["연찬_우선기관"]
    for row in ws.iter_rows(min_row=4, values_only=True):
        if isinstance(row[0], (int, float)) and row[1]:
            priority[norm_name(row[1])] = {"order": int(row[0]), "reason": row[2], "focus": row[3]}

    # 시범수집검증
    pilots = {}
    if "시범수집검증" in wb.sheetnames:
        ws = wb["시범수집검증"]
        header = None
        for row in ws.iter_rows(values_only=True):
            if row and row[0] == "기관":
                header = row
                continue
            if header and row and row[0] and row[0] != "종합":
                r = dict(zip(header, row))
                pilots[norm_name(r["기관"])] = {
                    "list": r.get("목록"), "detail": r.get("상세"), "attachment_check": r.get("첨부주소"),
                    "download": r.get("실제 다운로드"), "pdf": r.get("PDF 검증"),
                    "sample_title": r.get("검증용 게시물"), "published_at": r.get("게시일"),
                    "detail_url": r.get("상세 URL"), "final_verdict": str(r.get("최종 판정") or ""),
                    "checked_at": "2026-10-13",
                }

    ws = wb["전체_상세목록"]
    sources, seen_url, seen_org = [], {}, {}
    for row in ws.iter_rows(min_row=4, values_only=True):
        if not isinstance(row[0], (int, float)):
            continue
        no, cat, org, menu, link_type, url, prio, feat, kw, checked, note = (list(row) + [None] * 11)[:11]
        org = str(org).replace("\n", " ").strip()
        url = str(url or "").strip()
        sid = f"src-{int(no):03d}"
        if not url.startswith("https://"):
            errors.append(f"[자료원 {sid} {org}] URL 이 https:// 로 시작하지 않음: {url}")
        if url in seen_url:
            errors.append(f"[자료원 {sid} {org}] URL 중복 ({seen_url[url]})")
        if norm_name(org) in seen_org:
            errors.append(f"[자료원 {sid} {org}] 기관명 중복 ({seen_org[norm_name(org)]})")
        seen_url[url] = sid
        seen_org[norm_name(org)] = sid
        pilot = pilots.get(norm_name(org))
        readiness, basis = readiness_for(org, str(link_type or ""), pilot)
        p = priority.get(norm_name(org))
        item = {
            "source_id": sid,
            "audit_org": org,
            "source_group": source_group(str(cat)),
            "original_category": cat,
            "menu_name": menu,
            "source_url": url,
            "link_type": link_type,
            "learning_priority": prio,
            "learning_priority_member": bool(p),
            "learning_priority_order": p["order"] if p else None,
            "learning_priority_reason": p["reason"] if p else None,
            "collection_readiness": readiness,
            "collection_readiness_label": READINESS[readiness],
            "collection_readiness_basis": basis,
            "characteristics": feat,
            "recommended_keywords": [k.strip() for k in str(kw or "").split(",") if k.strip()],
            "source_checked_at": str(checked)[:10] if checked else None,
            "note": note,
        }
        if pilot:
            item["pilot_validation"] = pilot
        sources.append(item)

    missing = set(priority) - {norm_name(s["audit_org"]) for s in sources}
    for m in missing:
        errors.append(f"[연찬 우선기관] '{m}' 이 전체 목록에 없음")
    if len(sources) != 76:
        warnings.append(f"자료원 수가 76개가 아님: {len(sources)}")
    return {
        "schema_version": "2.0",
        "generated_at": now_kst(),
        "source_file": path.name,
        "readiness_definitions": READINESS,
        "counts": {
            "total": len(sources),
            "learning_priority_members": sum(s["learning_priority_member"] for s in sources),
            "by_readiness": {k: sum(1 for s in sources if s["collection_readiness"] == k) for k in READINESS},
        },
        "sources": sources,
    }


# ---------------------------------------------------------------- 유의어
def build_synonyms() -> dict:
    m = json.loads((ADMIN / "synonyms_master.json").read_text(encoding="utf-8"))
    blocked = {norm_name(t) for t in m["rules"]["blocked_generic_terms"]}
    groups = {}
    for key, g in m["groups"].items():
        terms = []
        for t in g["terms"]:
            if norm_name(t) in blocked:
                warnings.append(f"[유의어 {key}] 일반어 '{t}' 는 확장에서 제외됨")
                continue
            if t not in terms:
                terms.append(t)
        groups[key] = {"triggers": g["triggers"], "terms": terms}
    return {
        "schema_version": "2.0",
        "generated_at": now_kst(),
        "blocked_generic_terms": m["rules"]["blocked_generic_terms"],
        "groups": groups,
    }


# ---------------------------------------------------------------- 검토상태표
def write_review_sheet(cases: list[dict]) -> None:
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.worksheet.datavalidation import DataValidation

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "검토상태표"
    cols = [("사례번호", "id", 10), ("감사기관", "audit_org", 16), ("감사명", "audit_name", 26),
            ("게시일", "published_at", 11), ("분야", "category", 11), ("지적제목", "title", 44),
            ("요약", "summary", 50), ("처분", "disposition", 14), ("관련근거", "legal_basis", 18),
            ("검토상태", "review_status", 14), ("검토메모", "review_note", 36),
            ("검토일", "reviewed_at", 11), ("공식 게시물", "source_url", 40)]
    ws.append([c[0] for c in cols])
    for c in cases:
        ws.append([c.get(k) if c.get(k) is not None else "" for _, k, _ in cols])
    head = Font(bold=True, color="FFFFFF")
    for i, (_, _, w) in enumerate(cols, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font = head
        cell.fill = PatternFill("solid", fgColor="111111")
        ws.column_dimensions[cell.column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = ws.dimensions
    dv = DataValidation(type="list", formula1='"needs_review,source_checked,approved,rejected"', allow_blank=False)
    ws.add_data_validation(dv)
    dv.add(f"J2:J{len(cases) + 1}")

    g = wb.create_sheet("상태 설명")
    g.append(["검토상태", "화면 표시", "직원 화면 공개", "의미"])
    for k, v in REVIEW_STATUSES.items():
        g.append([k, v["label"], "공개" if v["public"] else "비공개", v["desc"]])
    g.append([])
    g.append(["현황", "건수"])
    for k in REVIEW_STATUSES:
        g.append([k, sum(1 for c in cases if c.get("review_status") == k)])
    for col, w in zip("ABCD", (16, 12, 14, 70)):
        g.column_dimensions[col].width = w
    g["A1"].font = g["B1"].font = g["C1"].font = g["D1"].font = Font(bold=True)
    wb.save(ADMIN / "review_status.xlsx")


# ---------------------------------------------------------------- 출력
def write_pair(name: str, key: str, data: dict) -> None:
    DIST_DATA.mkdir(parents=True, exist_ok=True)
    text = json.dumps(data, ensure_ascii=False, indent=2)
    (DIST_DATA / f"{name}.json").write_text(text + "\n", encoding="utf-8")
    js = ("/* 자동 생성 파일입니다. 직접 고치지 말고 admin/build_data.py 를 실행하세요. */\n"
          "window.__AUDIT_DATA__ = window.__AUDIT_DATA__ || {};\n"
          f"window.__AUDIT_DATA__.{key} = {text};\n")
    (DIST_DATA / f"{name}.js").write_text(js, encoding="utf-8")


def report() -> int:
    for w in warnings:
        print("[주의]", w)
    for e in errors:
        print("[오류]", e)
    return 1 if errors else 0


def main() -> int:
    cases_out, master_cases = build_cases()
    sources_out = build_sources()
    synonyms_out = build_synonyms()
    if errors:
        print("오류가 있어 파일을 만들지 않았습니다.")
        return report()
    write_pair("audit_cases", "cases", cases_out)
    write_pair("sources", "sources", sources_out)
    write_pair("synonyms", "synonyms", synonyms_out)
    write_review_sheet(master_cases)
    report()
    c, s = cases_out["counts"], sources_out["counts"]
    print(f"[완료] 감사사례: 전체 {c['master_total']}건 중 직원 화면 공개 {c['public_total']}건 {c['by_status']}")
    print(f"[완료] 감사자료원: {s['total']}개 (연찬 우선 {s['learning_priority_members']}곳) 수집준비도 {s['by_readiness']}")
    print(f"[완료] 유의어 그룹: {len(synonyms_out['groups'])}개")
    print("[완료] 검토상태표: admin/review_status.xlsx")
    return 0


if __name__ == "__main__":
    sys.exit(main())
